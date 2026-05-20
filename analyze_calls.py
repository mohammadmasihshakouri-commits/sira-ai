import os
import re
import time
import subprocess
from pathlib import Path

# -----------------------------
# Fix CUDA DLL paths on Windows
# -----------------------------

BASE_DIR = Path(__file__).resolve().parent

CUDA_DLL_PATHS = [
    BASE_DIR / "venv" / "Lib" / "site-packages" / "nvidia" / "cublas" / "bin",
    BASE_DIR / "venv" / "Lib" / "site-packages" / "nvidia" / "cudnn" / "bin",
    BASE_DIR / "venv" / "Lib" / "site-packages" / "nvidia" / "cuda_nvrtc" / "bin",
]

DLL_HANDLES = []

print("Checking CUDA DLL paths...")

for path in CUDA_DLL_PATHS:
    path_str = str(path)

    if path.exists():
        os.environ["PATH"] = path_str + os.pathsep + os.environ["PATH"]
        handle = os.add_dll_directory(path_str)
        DLL_HANDLES.append(handle)
        print("Added:", path_str)
    else:
        print("Not found:", path_str)

# -----------------------------
# Main imports
# -----------------------------

from faster_whisper import WhisperModel
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# -----------------------------
# Settings
# -----------------------------

SAMPLES_DIR = BASE_DIR / "call_samples"
CLEANED_DIR = BASE_DIR / "cleaned_calls"
OUTPUT_XLSX = BASE_DIR / "call_analysis.xlsx"

DEVICE = "cuda"
COMPUTE_TYPE = "float16"
WHISPER_MODEL_NAME = str(BASE_DIR / "models" / "faster-whisper-large-v3")

# فعلاً برای تست فقط ۵ فایل اول
# وقتی خروجی بهتر شد، این را None کن
TEST_LIMIT = 5

SUPPORTED_EXTENSIONS = {
    ".wav",
    ".mp3",
    ".m4a",
    ".ogg",
    ".flac",
    ".webm",
    ".aac",
    ".wma",
}


# -----------------------------
# Normalizer
# -----------------------------

def normalize_text_light(text):
    text = (text or "").strip()

    replacements = {
        # Greetings / courtesy
        "صبح تو، بخیر بشه": "صبحتون بخیر باشه",
        "صبح تو بخیر بشه": "صبحتون بخیر باشه",
        "صبح تو، بخیر باشه": "صبحتون بخیر باشه",
        "صبح تو بخیر باشه": "صبحتون بخیر باشه",
        "صبح تو بخیر": "صبحتون بخیر",
        "صبح تو": "صبحتون",
        "صبح تون": "صبحتون",
        "صبتون": "صبحتون",
        "صبتون بخیر": "صبحتون بخیر",
        "صبتون بخیر باشه": "صبحتون بخیر باشه",
        "صبحتون بخیر بشه": "صبحتون بخیر باشه",

        "ظهر تو": "ظهرتون",
        "ظهر تون": "ظهرتون",
        "ظهرتون بخیر بشه": "ظهرتون بخیر باشه",

        "عصر تو": "عصرتون",
        "عصر تون": "عصرتون",
        "عصرتون بخیر بشه": "عصرتون بخیر باشه",

        "شب تو": "شبتون",
        "شب تون": "شبتون",
        "شبتون بخیر بشه": "شبتون بخیر باشه",

        "وقت تو": "وقتتون",
        "وقت تون": "وقتتون",
        "وقتون": "وقتتون",
        "وفتتون": "وقتتون",
        "فقتتون": "وقتتون",
        "فقتون": "وقتتون",
        "بقتتون": "وقتتون",
        "بقتون": "وقتتون",
        "وقتتون خیر": "وقتتون بخیر",
        "وقتتون بخیر بشه": "وقتتون بخیر باشه",

        "بخیر بشه": "بخیر باشه",
        "بخیرش": "بخیر باشه",
        "به خیر": "بخیر",
        "بحیر": "بخیر",

        "سلام جنب": "سلام جناب",
        "سلام جنم": "سلام جناب",
        "سلام جانب": "سلام جناب",
        "سلام جانبی": "سلام جناب",
        "سلام جنابی": "سلام جناب",
        "جنب": "جناب",
        "جنم": "جناب",
        "جنابی": "جناب",

        "خسته باشید": "خسته نباشید",
        "خسته نوشید": "خسته نباشید",
        "خسته توشید": "خسته نباشید",
        "خسته نبشید": "خسته نباشید",

        # Ticket / payment / SMS
        "بلیط": "بلیت",
        "بلیطم": "بلیتم",
        "بیلیت": "بلیت",
        "بیریت": "بلیت",
        "خرید بلیته": "بلیت خریده",
        "خرید بلیتاری": "بلیت خریداری",
        "اس ام اس": "پیامک",
        "اسمس": "پیامک",
        "SMS": "پیامک",
        "پرداختشم انجام دادم": "پرداخت انجام دادم",
        "پرداختشم": "پرداخت",

        # Cancellation wording
        "کانسلش کنم": "کنسلش کنم",
        "کانسلش": "کنسلش",
        "کانسل کنم": "کنسل کنم",
        "کانسل": "کنسل",
        "کنصل": "کنسل",
        "می‌خوام پسش بدن": "می‌خوام پسش بدم",
        "میخوام پسش بدن": "می‌خوام پسش بدم",
        "پسش بدن": "پسش بدم",
        "پس بدن": "پس بدم",

        # Numbers
        "صرف": "صفر",
        "سفر": "صفر",
        "نومصد": "نهصد",
        "نومسد": "نهصد",
        "نصد": "نهصد",
        "نوهصد": "نهصد",
        "دویصد": "دویست",
        "چارصد": "چهارصد",
        "چل": "چهل",
        "شست": "شصت",
        "شیش": "شش",
        "پنجا": "پنجاه",
        "پونصد": "پانصد",
        "بیستو دو": "بیست و دو",
        "بیستودو": "بیست و دو",
        "سیاسه": "سی و سه",

        # Common words
        "اندوز": "هنوز",
        "انوز": "هنوز",
        "نمی‌مده": "نیومده",
        "نمیومده": "نیومده",
    }

    for wrong, correct in replacements.items():
        text = text.replace(wrong, correct)

    text = re.sub(r"\s+", " ", text).strip()
    return text


# -----------------------------
# Helpers
# -----------------------------

def extract_simple_digits(text):
    if not text:
        return ""

    ascii_text = text.translate(
        str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    )

    digits = "".join(re.findall(r"\d", ascii_text))

    if 4 <= len(digits) <= 20:
        return digits

    return ""


def is_repetitive_or_hallucinated(text):
    text = normalize_text_light(text)
    tokens = re.findall(r"[\w\u0600-\u06FF]+", text)

    if len(tokens) < 6:
        return False

    token_counts = {}

    for token in tokens:
        token_counts[token] = token_counts.get(token, 0) + 1

    most_repeated = max(token_counts.values()) if token_counts else 0

    if most_repeated >= 8:
        return True

    repeated_bad_phrases = [
        "شبتون بخیر",
        "بخیر",
        "بلیت",
        "کنسلی",
        "مرکز تماس",
        "مرکز تماسیم",
        "سلامتیم",
    ]

    for phrase in repeated_bad_phrases:
        if text.count(phrase) >= 5:
            return True

    return False


def looks_like_bad_stt(text):
    text = normalize_text_light(text)

    if not text or len(text.strip()) < 4:
        return True

    tokens = re.findall(r"[\w\u0600-\u06FF]+", text)

    if len(tokens) <= 1:
        return True

    if len(tokens) >= 8:
        unique_ratio = len(set(tokens)) / len(tokens)

        if unique_ratio < 0.35:
            return True

    if is_repetitive_or_hallucinated(text):
        return True

    return False


def suggest_intent(text):
    text = normalize_text_light(text)

    cancellation_keywords = [
        "کنسل",
        "کنسلی",
        "استرداد",
        "لغو",
        "پسش بدم",
        "پس بدم",
        "پولم برگرده",
        "عودت",
        "مرجوع",
    ]

    sms_keywords = [
        "پیامک",
        "نیومده",
        "نرسیده",
        "ارسال نشده",
        "دریافت نکردم",
    ]

    payment_keywords = [
        "پرداخت",
        "پول",
        "کم شده",
        "کسر شده",
        "تراکنش",
        "حساب",
    ]

    reservation_code_keywords = [
        "کد رزرو",
        "کد کجاست",
        "بارکد",
        "جزئیات بلیت",
        "بلیت‌های من",
    ]

    tiwall_keywords = [
        "تیوال",
        "تی وال",
    ]

    audio_keywords = [
        "صداتون",
        "صدا",
        "نمی‌شنوم",
        "واضح نیست",
        "قطع شد",
        "تکرار",
        "دوباره",
    ]

    if any(k in text for k in tiwall_keywords):
        return "tiwall"

    if any(k in text for k in cancellation_keywords):
        return "cancellation"

    if any(k in text for k in reservation_code_keywords):
        return "reservation_code_help"

    if any(k in text for k in payment_keywords) and any(k in text for k in sms_keywords):
        return "payment_sms_missing"

    if any(k in text for k in sms_keywords):
        return "ticket_sms_missing"

    if any(k in text for k in payment_keywords):
        return "payment_issue"

    if any(k in text for k in audio_keywords):
        return "audio_issue"

    if (
        "سلام" in text
        or "وقتتون بخیر" in text
        or "صبحتون بخیر" in text
        or "ظهرتون بخیر" in text
        or "عصرتون بخیر" in text
        or "شبتون بخیر" in text
        or "خسته نباشید" in text
    ):
        return "greeting"

    return "other"


def quality_flag(raw_text, normalized_text, detected_number):
    if not raw_text.strip():
        return "bad_stt_empty"

    if looks_like_bad_stt(normalized_text):
        return "bad_stt_needs_review"

    if len(normalized_text) < 10:
        return "too_short_needs_review"

    if any(word in normalized_text for word in ["کد", "شماره", "موبایل", "رزرو"]) and not detected_number:
        return "number_uncertain_needs_review"

    return "good"


def save_rows_to_xlsx(rows, fieldnames):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "call_analysis"

    sheet.append(fieldnames)

    for row in rows:
        sheet.append([row.get(field, "") for field in fieldnames])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wrap_columns = {
        "raw_text",
        "normalized_text",
        "manual_correction",
        "final_text",
        "notes",
        "segments",
    }

    for col_index, field_name in enumerate(fieldnames, start=1):
        column_letter = get_column_letter(col_index)

        if field_name in wrap_columns:
            sheet.column_dimensions[column_letter].width = 45
        else:
            sheet.column_dimensions[column_letter].width = 20

        for cell in sheet[column_letter]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    workbook.save(OUTPUT_XLSX)
def clean_audio_with_ffmpeg(input_path):
    """
    Converts noisy call-center audio to a cleaner 16kHz mono WAV file.
    This is for offline analysis only, not real-time demo.
    """

    CLEANED_DIR.mkdir(exist_ok=True)

    output_path = CLEANED_DIR / f"{input_path.stem}_clean.wav"

    if output_path.exists():
        return output_path

    command = [
        "ffmpeg",
        "-y",
        "-i", str(input_path),

        # Convert to mono 16k wav and apply light phone-call cleanup
        "-ac", "1",
        "-ar", "16000",
        "-af",
        "highpass=f=120,lowpass=f=3800,afftdn=nf=-25,loudnorm=I=-18:TP=-1.5:LRA=11",

        str(output_path),
    ]

    subprocess.run(
        command,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=True,
    )

    return output_path

# -----------------------------
# Main
# -----------------------------

def main():
    if not SAMPLES_DIR.exists():
        print(f"Folder not found: {SAMPLES_DIR}")
        print("Please create call_samples folder and put audio files inside it.")
        return

    audio_files = [
        path for path in sorted(SAMPLES_DIR.iterdir())
        if path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    if TEST_LIMIT is not None:
        audio_files = audio_files[:TEST_LIMIT]

    if not audio_files:
        print(f"No audio files found in: {SAMPLES_DIR}")
        return

    print("Loading Whisper model...")
    print(f"Model: {WHISPER_MODEL_NAME} | Device: {DEVICE} | Compute: {COMPUTE_TYPE}")

    model = WhisperModel(
        WHISPER_MODEL_NAME,
        device=DEVICE,
        compute_type=COMPUTE_TYPE,
        cpu_threads=8,
        num_workers=1,
    )

    print(f"Found {len(audio_files)} audio files.")
    print("Starting transcription...")

    rows = []

    for index, audio_path in enumerate(audio_files, start=1):
        print(f"[{index}/{len(audio_files)}] Processing: {audio_path.name}")

        start_time = time.time()

        try:
            cleaned_audio_path = clean_audio_with_ffmpeg(audio_path)
            
            segments, info = model.transcribe(
                str(cleaned_audio_path),
                language="fa",
                task="transcribe",
                beam_size=1,
                best_of=1,
                condition_on_previous_text=False,
                vad_filter=True,
                vad_parameters=dict(
                    threshold=0.8,
                    min_speech_duration_ms=700,
                    min_silence_duration_ms=900,
                    speech_pad_ms=150,
                ),
                initial_prompt=None,
                temperature=0.0,
                compression_ratio_threshold=2.0,
                log_prob_threshold=-0.6,
                no_speech_threshold=0.75,
            )

            raw_parts = []
            segment_parts = []

            for segment in segments:
                cleaned = segment.text.strip()

                if not cleaned:
                    continue

                if is_repetitive_or_hallucinated(cleaned):
                    segment_parts.append(
                        f"[{round(segment.start, 2)}-{round(segment.end, 2)}] SKIPPED_REPETITIVE: {cleaned}"
                    )
                    continue

                raw_parts.append(cleaned)
                segment_parts.append(
                    f"[{round(segment.start, 2)}-{round(segment.end, 2)}] {cleaned}"
                )

            raw_text = " ".join(raw_parts).strip()
            normalized_text = normalize_text_light(raw_text)
            detected_number = extract_simple_digits(normalized_text)
            intent = suggest_intent(normalized_text)
            q_flag = quality_flag(raw_text, normalized_text, detected_number)

            stt_time = round(time.time() - start_time, 2)
            language = getattr(info, "language", "")
            language_probability = round(getattr(info, "language_probability", 0), 2)
            duration = round(getattr(info, "duration", 0), 2) if hasattr(info, "duration") else ""

            rows.append({
                "file_name": audio_path.name,
                "raw_text": raw_text,
                "normalized_text": normalized_text,
                "detected_number": detected_number,
                "suggested_intent": intent,
                "quality_flag": q_flag,
                "manual_correction": "",
                "final_text": "",
                "notes": "",
                "language": language,
                "language_probability": language_probability,
                "duration_seconds": duration,
                "stt_time_seconds": stt_time,
                "segments": " | ".join(segment_parts),
            })

        except Exception as e:
            stt_time = round(time.time() - start_time, 2)

            rows.append({
                "file_name": audio_path.name,
                "raw_text": "",
                "normalized_text": "",
                "detected_number": "",
                "suggested_intent": "error",
                "quality_flag": "error",
                "manual_correction": "",
                "final_text": "",
                "notes": str(e),
                "language": "",
                "language_probability": "",
                "duration_seconds": "",
                "stt_time_seconds": stt_time,
                "segments": "",
            })

            print(f"Error processing {audio_path.name}: {e}")

    fieldnames = [
        "file_name",
        "raw_text",
        "normalized_text",
        "detected_number",
        "suggested_intent",
        "quality_flag",
        "manual_correction",
        "final_text",
        "notes",
        "language",
        "language_probability",
        "duration_seconds",
        "stt_time_seconds",
        "segments",
    ]

    save_rows_to_xlsx(rows, fieldnames)

    print("")
    print("Done.")
    print(f"Output saved to: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()